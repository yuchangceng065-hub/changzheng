"""把 docx 的段落与表格内容导出为纯文本，便于比对线路内容。

用法：
    python scripts/dump-docx.py <input.docx> [output.txt]
"""

import sys
from docx import Document


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: dump-docx.py <input.docx> [output.txt]", file=sys.stderr)
        return 2

    doc = Document(sys.argv[1])
    lines: list[str] = []

    lines.append("=== PARAGRAPHS ===")
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            style = para.style.name if para.style is not None else ""
            lines.append(f"[{i}][{style}] {text}")

    lines.append("")
    lines.append(f"=== TABLES ({len(doc.tables)}) ===")
    for ti, table in enumerate(doc.tables):
        lines.append(f"--- table {ti} ({len(table.rows)} rows) ---")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            lines.append(" | ".join(cells))

    out = "\n".join(lines)

    if len(sys.argv) >= 3:
        with open(sys.argv[2], "w", encoding="utf-8") as fh:
            fh.write(out)
        print(f"wrote {sys.argv[2]} ({len(lines)} lines)")
    else:
        print(out)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
